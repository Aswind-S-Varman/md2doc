from docx import Document

from md2doc import python_backend, converter


def test_separator_row_is_detected():
    assert python_backend._is_separator(["---", "---"])
    assert not python_backend._is_separator(["Name", "Role"])


def test_table_rows_exclude_separator():
    lines = [
        "| Name | Role |",
        "|------|------|",
        "| Alice | Engineer |",
    ]

    rows, next_index = python_backend._collect_table(lines, 0)

    assert rows == [["Name", "Role"], ["Alice", "Engineer"]]
    assert next_index == 3


def test_python_backend_writes_heading_and_table(tmp_path):
    source = tmp_path / "doc.md"
    source.write_text(
        "# Title\n\n| Name | Role |\n|------|------|\n| Alice | Engineer |\n",
        encoding="utf-8",
    )
    target = tmp_path / "doc.docx"

    python_backend.convert(source, target)

    result = Document(target)
    assert result.paragraphs[0].text == "Title"
    assert len(result.tables) == 1
    assert result.tables[0].cell(1, 0).text == "Alice"


def test_bullet_list_uses_bullet_style(tmp_path):
    source = tmp_path / "doc.md"
    source.write_text("- First\n- Second\n", encoding="utf-8")
    target = tmp_path / "doc.docx"

    python_backend.convert(source, target)

    result = Document(target)
    assert [p.text for p in result.paragraphs] == ["First", "Second"]
    assert all(p.style.name == "List Bullet" for p in result.paragraphs)


def test_converter_reports_backend_used(tmp_path):
    source = tmp_path / "doc.md"
    source.write_text("# Hello\n", encoding="utf-8")
    target = tmp_path / "doc.docx"

    used = converter.convert(source, target, backend="python")

    assert used == "python"
    assert target.exists()


