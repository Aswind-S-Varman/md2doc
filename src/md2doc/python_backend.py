from pathlib import Path

from docx import Document


def convert(source: Path, target: Path) -> None:
    doc = Document()
    lines = source.read_text(encoding="utf-8").splitlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            i += 1
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 9))
            i += 1
        elif line.lstrip().startswith("|"):
            rows, i = _collect_table(lines, i)
            _add_table(doc, rows)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
            i += 1
        else:
            doc.add_paragraph(line)
            i += 1

    doc.save(target)


def _collect_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start

    while i < len(lines) and lines[i].lstrip().startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not _is_separator(cells):
            rows.append(cells)
        i += 1

    return rows, i


def _is_separator(cells: list[str]) -> bool:
    """Markdown's |---|---| alignment row carries no data."""
    return all(cell and set(cell) <= set("-: ") for cell in cells)


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return

    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"

    for r, row in enumerate(rows):
        for c, value in enumerate(row[: len(rows[0])]):
            table.cell(r, c).text = value


